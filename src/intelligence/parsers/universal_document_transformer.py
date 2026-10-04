"""
ROXIRA UNIVERSAL DOCUMENT CONVERSION & EXPORT ENGINE
مرکز تبدیل و صدور اسناد به PDF، Word، Markdown و JSON
"""
import sys
import json
import pathlib
from typing import Dict, Any
from docx import Document

CURRENT_DIR = pathlib.Path.cwd()
if str(CURRENT_DIR) not in sys.path:
    sys.path.insert(0, str(CURRENT_DIR))

class UniversalDocumentTransformer:
    @staticmethod
    def to_structured_json(raw_text: str, title: str = "سند سازمانی") -> Dict[str, Any]:
        """تبدیل متن خام به ساختار استاندارد JSON"""
        lines = [l.strip() for l in raw_text.splitlines() if l.strip()]
        sections = []
        current_section = {"heading": "مقدمه", "content": []}
        
        for line in lines:
            if line.startswith("ماده") or line.startswith("#") or line.startswith("بخش"):
                if current_section["content"]:
                    sections.append(current_section)
                current_section = {"heading": line.lstrip("#").strip(), "content": []}
            else:
                current_section["content"].append(line)
        if current_section["content"]:
            sections.append(current_section)

        return {
            "document_title": title,
            "total_sections": len(sections),
            "sections": sections
        }

    @staticmethod
    def to_markdown(raw_text: str, title: str = "سند سازمانی") -> str:
        """تبدیل متن به فرمت تمیز Markdown"""
        lines = [l.strip() for l in raw_text.splitlines() if l.strip()]
        md_lines = [f"# {title}\n"]
        for line in lines:
            if line.startswith("ماده") or line.startswith("فصل"):
                md_lines.append(f"\n### {line}")
            elif ":" in line and len(line.split(":", 1)[0]) < 25:
                k, _, v = line.partition(":")
                md_lines.append(f"* **{k.strip()}**: {v.strip()}")
            else:
                md_lines.append(line)
        return "\n".join(md_lines)

    @staticmethod
    def export_to_docx(title: str, body_text: str, output_path: str) -> str:
        """تولید فایل رسمی Word (.docx)"""
        doc = Document()
        doc.add_heading(title, level=1)
        for p in body_text.splitlines():
            if p.strip():
                doc.add_paragraph(p.strip())
        doc.save(output_path)
        return output_path

    @staticmethod
    def export_to_pdf_html(title: str, body_text: str, output_path: str) -> str:
        """تولید قالب چاپی استاندارد A4 قابل ذخیره به صورت PDF"""
        html_content = f"""<!DOCTYPE html>
<html lang="fa" dir="rtl">
<head>
    <meta charset="UTF-8">
    <title>{title}</title>
    <style>
        body {{ font-family: 'Tahoma', sans-serif; padding: 40px; color: #1e293b; }}
        h1 {{ border-bottom: 2px solid #cbd5e1; padding-bottom: 10px; font-size: 20px; text-align: center; }}
        p {{ line-height: 1.8; font-size: 14px; text-align: justify; }}
        .meta {{ font-size: 12px; color: #64748b; margin-top: 30px; border-top: 1px dashed #cbd5e1; padding-top: 10px; }}
    </style>
</head>
<body>
    <h1>{title}</h1>
    {"".join(f"<p>{p.strip()}</p>" for p in body_text.splitlines() if p.strip())}
    <div class="meta">تولیدشده توسط پلتفرم هوشمند ROXIRA | سند رسمی دیجیتال</div>
</body>
</html>"""
        pathlib.Path(output_path).write_text(html_content, encoding="utf-8")
        return output_path

if __name__ == "__main__":
    transformer = UniversalDocumentTransformer()
    sample_text = "ماده ۲۳ قانون صدور چک\nدارنده چک با ارائه گواهی عدم پرداخت تقاضای اجرائیه مستقیم می‌نماید.\nخواهان: ماهان محمدی\nمبلغ: ۱۰۰,۰۰۰,۰۰۰ تومان"
    
    json_out = transformer.to_structured_json(sample_text, "دادخواست چک")
    docx_path = transformer.export_to_docx("دادخواست صدور اجرائیه", sample_text, "output_apps/smart_judge/sample_export.docx")
    pdf_html_path = transformer.export_to_pdf_html("دادخواست صدور اجرائیه", sample_text, "output_apps/smart_judge/sample_export.html")
    
    print("✅ Document Transformer Tested:")
    print("  ├─ JSON Sections:", json_out["total_sections"])
    print("  ├─ DOCX Export:", docx_path)
    print("  └─ Printable/PDF HTML:", pdf_html_path)
