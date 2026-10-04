import sys
import pathlib
from docx import Document

class OfficialDocxExporter:
    @staticmethod
    def create_lawsuit_docx(draft_data: dict, output_filename: str = "official_lawsuit.docx") -> str:
        doc = Document()
        doc.add_heading(draft_data.get("lawsuit_title", "دادخواست قضایی"), level=1)

        # مشخصات طرفین
        doc.add_heading("مشخصات طرفین و خواسته:", level=2)
        doc.add_paragraph(f"خواهان: {draft_data.get('plaintiff', '---')}")
        doc.add_paragraph(f"خوانده: {draft_data.get('defendant', '---')}")
        if draft_data.get("claim_amount_toman"):
            doc.add_paragraph(f"بهای خواسته: {draft_data.get('claim_amount_toman'):,.0f} تومان")

        # مواد قانونی
        doc.add_heading("دلایل و مستندات قانونی:", level=2)
        for stat in draft_data.get("statutes_cited", []):
            doc.add_paragraph(f"• {stat}")

        # متن دادخواست
        doc.add_heading("شرح دادخواست:", level=2)
        doc.add_paragraph(draft_data.get("petition_text", ""))

        out_path = pathlib.Path("output_apps/smart_judge") / output_filename
        doc.save(str(out_path))
        return str(out_path)

if __name__ == "__main__":
    sample_data = {
        "lawsuit_title": "درخواست صدور اجرائیه چک صیادی",
        "plaintiff": "ماهان محمدی",
        "defendant": "صادرکننده چک",
        "claim_amount_toman": 100000000.0,
        "statutes_cited": ["ماده ۲۳ قانون صدور چک"],
        "petition_text": "تقاضای صدور اجرائیه مستقیم مطابق گواهی عدم پرداخت."
    }
    path = OfficialDocxExporter.create_lawsuit_docx(sample_data)
    print("✅ Official DOCX Exporter Tested:", path)
